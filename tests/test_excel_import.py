from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import traceback
import pandas as pd
import pytest
from openpyxl import Workbook
from analytics import (load_grid, list_sheets, report_from_grid, detect_header,
                       read_report, normalize, infer_mapping)


def workbook_bytes():
    wb = Workbook()
    wb.active.title = '说明'
    wb.active.append(['不是订单数据'])
    ws = wb.create_sheet('备货管理')
    ws.append(['Temu备货导出'])
    ws.merge_cells('A1:G1')
    ws.append([])
    ws.append(['订单日期', 'SKU\u200b', '采购数量', '供货单价', '订单状态', '订单号', '条码'])
    ws.append(['2026年10月01日', '001_ABC', '2', '￥3.50', '已完成', '000_01', '000123'])
    ws.append([])
    ws.append(['2026-10-02T08:00:00+08:00', 'NA', '1', '2', '待发货', '000_02', '001234'])
    ws['H4'] = 999999
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def malformed_numeric():
    original = workbook_bytes()
    out = BytesIO()
    with ZipFile(BytesIO(original)) as src, ZipFile(out, 'w', ZIP_DEFLATED) as dst:
        for name in src.namelist():
            content = src.read(name)
            if name == 'xl/worksheets/sheet2.xml':
                content = content.replace(b'<v>999999</v>', b'<v>_</v>')
            dst.writestr(name, content)
    return out.getvalue()


def test_reproduce_exact_openpyxl_error_then_recover():
    data = malformed_numeric()
    with pytest.raises(ValueError, match="invalid literal for int.*'_'"):
        pd.read_excel(BytesIO(data), sheet_name='备货管理', dtype=str)
    try:
        pd.read_excel(BytesIO(data), sheet_name='备货管理', dtype=str)
    except ValueError:
        assert '_cast_number' in traceback.format_exc()
    assert list_sheets(data, 'report.xlsx') == ['说明', '备货管理']
    grid = load_grid(data, 'report.xlsx', '备货管理')
    assert grid.iat[3, 7] == '_'
    assert any('H4' in v and "'_'" in v for v in grid.attrs['diagnostics'])
    assert any('_cast_number' in v for v in grid.attrs['diagnostics'])
    report = report_from_grid(grid, detect_header(grid))
    valid, rejected = normalize(report, infer_mapping(report.columns))
    assert rejected.empty
    assert valid['SKU'].tolist() == ['001_ABC', 'NA']
    assert valid['订单号'].tolist() == ['000_01', '000_02']
    assert valid['条码'].tolist() == ['000123', '001234']
    assert valid['源数据行'].tolist() == [4, 6]
    assert valid['采购金额'].sum() == 9


def test_offset_header_merges_empty_rows_and_text():
    grid = load_grid(workbook_bytes(), 'report.xlsx', '备货管理')
    assert detect_header(grid) == 2
    raw = report_from_grid(grid, 2)
    assert len(raw) == 2
    assert raw.iloc[1]['SKU\u200b'] == 'NA'
    valid, bad = normalize(raw, infer_mapping(raw.columns))
    assert bad.empty
    assert valid['源数据行'].tolist() == [4, 6]
    assert valid.iloc[0]['日期'] == pd.Timestamp('2026-10-01')


def test_duplicate_blank_headers_and_invalid_numeric_retained():
    grid = pd.DataFrame([['订单日期', 'SKU', '采购数量', '供货价', '订单状态', 'SKU', None],
                         ['46023', '00001', '_', '1', '完成', 'different', 'NA']])
    raw = report_from_grid(grid)
    assert raw.columns.is_unique
    assert 'SKU [2]' in raw.columns
    assert '未命名列7' in raw.columns
    valid, rejected = normalize(raw, infer_mapping(raw.columns))
    assert valid.empty and len(rejected) == 1
    assert rejected.iloc[0]['采购数量'] == '_'
    assert '数量必须为正整数' in rejected.iloc[0]['无效原因']


def test_1904_epoch_fallback():
    from openpyxl.utils.datetime import CALENDAR_MAC_1904
    wb = Workbook()
    wb.epoch = CALENDAR_MAC_1904
    ws = wb.active
    ws.append(['订单日期', 'SKU', '采购数量', '供货价', '订单状态', '其他'])
    ws.append([44561, '001', 1, 2, '完成', 999999])
    buf = BytesIO()
    wb.save(buf)
    out = BytesIO()
    with ZipFile(BytesIO(buf.getvalue())) as src, ZipFile(out, 'w') as dst:
        for n in src.namelist():
            dst.writestr(n, src.read(n).replace(b'<v>999999</v>', b'<v>_</v>'))
    raw = read_report(out.getvalue(), 'date.xlsx')
    valid, bad = normalize(raw, infer_mapping(raw.columns))
    assert bad.empty
    assert valid.iloc[0]['日期'] == pd.Timestamp('2026-01-01')


def test_multilevel_headers_and_local_calendar_day():
    grid = pd.DataFrame([
        ['订单信息', None, '采购信息', None, '订单信息'],
        ['订单日期', 'SKU', '采购数量', '供货价', '订单状态'],
        ['2026-10-02T00:30:00+08:00', '001', '1', '2', '完成'],
    ])
    raw = report_from_grid(grid, 0, 2)
    valid, rejected = normalize(raw, infer_mapping(raw.columns))
    assert rejected.empty
    assert valid.iloc[0]['日期'] == pd.Timestamp('2026-10-02')
    assert valid.iloc[0]['源数据行'] == 3


def test_style_parse_failure_is_diagnosed():
    # A malformed integer in styles can fail before worksheet parsing.
    source = workbook_bytes()
    out = BytesIO()
    with ZipFile(BytesIO(source)) as src, ZipFile(out, 'w') as dst:
        for name in src.namelist():
            content = src.read(name)
            if name == 'xl/styles.xml':
                content = content.replace(b'family val="2"', b'family val="_"')
            dst.writestr(name, content)
    grid = load_grid(out.getvalue(), 'style.xlsx', '备货管理')
    assert grid.attrs['diagnostics']
    assert "could not convert string to float" in grid.attrs['diagnostics'][1]
    valid, bad = normalize(report_from_grid(grid, 2), infer_mapping(report_from_grid(grid, 2).columns))
    assert len(valid) == 2 and bad.empty


def test_corrupt_file_has_explicit_failure():
    from zipfile import BadZipFile
    with pytest.raises(BadZipFile):
        list_sheets(b'not an Excel file', 'broken.xlsx')


def test_uploaded_malformed_excel_ui(monkeypatch):
    from streamlit.testing.v1 import AppTest
    import streamlit
    class Upload(BytesIO):
        name = 'synthetic.xlsx'
    monkeypatch.setattr(streamlit, 'file_uploader', lambda *args, **kwargs: Upload(malformed_numeric()))
    app = AppTest.from_file('app.py', default_timeout=30).run()
    next(r for r in app.radio if r.label == '选择数据').set_value('上传报表').run()
    next(s for s in app.selectbox if s.label == '工作表').set_value('备货管理').run()
    assert not app.exception
    assert [m.value for m in app.metric] == ['2', '3', '9.00']
    assert len(app.get('plotly_chart')) == 2
    assert len(app.get('download_button')) == 2
    assert next(n for n in app.number_input if n.label.startswith('表头所在行')).value == 3
    assert any('_cast_number' in c.value for c in app.code)


def test_corrupt_upload_ui_does_not_crash(monkeypatch):
    from streamlit.testing.v1 import AppTest
    import streamlit
    class Upload(BytesIO):
        name = 'broken.xlsx'
    monkeypatch.setattr(streamlit, 'file_uploader', lambda *args, **kwargs: Upload(b'broken'))
    app = AppTest.from_file('app.py', default_timeout=30).run()
    next(r for r in app.radio if r.label == '选择数据').set_value('上传报表').run()
    assert not app.exception
    assert any('BadZipFile' in e.value for e in app.error)
    assert any('Traceback' in c.value for c in app.code)


def test_only_actual_merged_text_can_be_filled():
    from analytics import fill_merged_text
    wb = Workbook()
    ws = wb.active
    ws.append(['订单日期', 'SKU', '数量', '供货价', '订单状态', '订单号'])
    ws.append(['2026-01-01', '001', 1, 2, '完成', '000_A'])
    ws.append(['2026-01-01', '002', 3, 2, '完成', None])
    ws.append(['2026-01-01', '003', 1, 2, '完成', None])
    ws.merge_cells('F2:F3')
    ws.merge_cells('D2:D3')  # Numeric merges must not be filled.
    buf = BytesIO()
    wb.save(buf)
    raw = read_report(buf.getvalue(), 'merged.xlsx')
    filled = fill_merged_text(raw, infer_mapping(raw.columns))
    assert filled.iloc[1]['订单号'] == '000_A'
    assert pd.isna(filled.iloc[2]['订单号'])
    assert pd.isna(filled.iloc[1]['供货价'])
    assert pd.isna(raw.iloc[1]['订单号'])
