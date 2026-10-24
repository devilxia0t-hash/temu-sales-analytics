from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from analytics import *


def fixture():
    return pd.DataFrame({'订单日期': ['2026-01-01', '2026-01-01', '2026-01-03', 'bad'],
                         'SKU': ['001', '002', '001', '003'], '采购数量': ['2', '3', '4', '-1'],
                         '供货单价': ['￥1,000', '2.50', '0', '5'], '订单状态': ['完成']*4,
                         '订单号': ['a', 'a', 'b', 'c']})


def test_calculation_and_rejection():
    raw = fixture()
    valid, rejected = normalize(raw, infer_mapping(raw.columns))
    assert len(valid) == 3 and len(rejected) == 1
    assert valid['采购金额'].sum() == 2007.5
    daily = daily_summary(valid, '2026-01-01', '2026-01-03', True)
    assert daily['采购订单数'].tolist() == [1, 0, 1]
    assert daily['采购件数'].tolist() == [5, 0, 4]
    assert ranking(valid, 'SKU').iloc[0]['SKU'] == '001'
    assert daily_summary(valid, '2026-01-01', '2026-01-03', False).iloc[0, 0] == 2


def test_invalid_and_serial_dates():
    raw = fixture()
    raw.loc[0, '订单日期'] = '46023'
    raw.loc[1, '采购数量'] = '1.5'
    raw.loc[2, '供货单价'] = 'inf'
    valid, rejected = normalize(raw, infer_mapping(raw.columns))
    assert len(valid) == 1 and len(rejected) == 3
    assert valid.iloc[0]['日期'] == pd.Timestamp('2026-01-01')


def test_read_formats_and_identifiers():
    raw = fixture()
    for encoding in ['utf-8-sig', 'gb18030']:
        result = read_report(raw.to_csv(index=False).encode(encoding), 'test.csv')
        assert result.iloc[0]['SKU'] == '001'
    buf = BytesIO()
    raw.to_excel(buf, index=False)
    assert read_report(buf.getvalue(), 'test.xlsx').iloc[0]['SKU'] == '001'


def test_growth_anomaly_and_export():
    daily = pd.DataFrame({'采购件数': [10]*59+[100], '采购金额': [1]*60}, index=pd.date_range('2026-01-01', periods=60, name='日期'))
    changes, anomalies = signals(daily)
    assert len(anomalies) == 1
    assert changes.iloc[0]['本期采购件数'] == 160
    assert abs(changes.iloc[0]['增长率'] - (160/70-1)) < 1e-9
    short, _ = signals(daily.head(5))
    assert short['增长率'].isna().all()
    raw = fixture()
    raw.loc[0, 'SKU'] = '=1+1'
    valid, rejected = normalize(raw, infer_mapping(raw.columns))
    blob = export_excel(valid, daily, changes, anomalies, rejected, [('口径', '采购非销售')])
    wb = load_workbook(BytesIO(blob))
    assert len(wb.sheetnames) == 10
    cell = next(c for row in wb['有效采购明细'] for c in row if c.value == '=1+1')
    assert cell.data_type == 's'


def test_streamlit_demo_and_controls():
    app = AppTest.from_file('app.py', default_timeout=30).run()
    assert not app.exception
    assert len(app.metric) == 3
    assert app.metric[0].label == '采购订单数'
    next(r for r in app.radio if r.label.startswith('趋势窗口')).set_value(30).run()
    assert not app.exception
    next(r for r in app.radio if r.label == '排行维度').set_value('尺码').run()
    assert not app.exception
    app.multiselect[0].set_value([]).run()
    assert not app.exception
    assert any('没有有效采购记录' in x.value for x in app.info)


def test_manual_mapping_and_missing_columns():
    import pytest
    raw = fixture().rename(columns={'SKU': '商家自定义编码'})
    mapping = infer_mapping(raw.columns)
    with pytest.raises(ValueError, match='SKU'):
        normalize(raw, mapping)
    mapping['SKU'] = '商家自定义编码'
    valid, _ = normalize(raw, mapping)
    assert len(valid) == 3
    mapping['SPU'] = '商家自定义编码'
    with pytest.raises(ValueError, match='多个字段'):
        normalize(raw, mapping)


def test_zero_baseline_and_status_filter():
    values = pd.DataFrame({'采购件数': [0]*7 + [10]*7}, index=pd.date_range('2026-01-01', periods=14))
    changes, _ = signals(values)
    assert pd.isna(changes.iloc[0]['增长率'])
    assert '上期为0' in changes.iloc[0]['说明']
    app = AppTest.from_file('app.py', default_timeout=30).run()
    app.multiselect[0].set_value(['已取消']).run()
    assert not app.exception
    assert int(app.metric[0].value.replace(',', '')) > 0
    assert app.metric[0].value != AppTest.from_file('app.py').run().metric[0].value
